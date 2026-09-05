import io
import logging
from typing import List, Dict

logger = logging.getLogger(__name__)

class ParsingError(Exception):
    """Custom exception raised when document parsing fails."""
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message

def extract_text(file_bytes: bytes, content_type: str, filename: str) -> List[Dict[str, str]]:
    """
    Extracts text blocks from PDF, DOCX, or PPTX bytes.
    Returns a list of dicts: [{"text": "...", "sourceRef": "Page 1"}]
    Raises ParsingError if file parsing fails.
    """
    fn_lower = filename.lower()
    blocks: List[Dict[str, str]] = []

    try:
        # 1. PDF Parsing
        if fn_lower.endswith('.pdf') or 'pdf' in content_type:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                if not pdf.pages:
                    raise ParsingError("PDF file contains no readable pages.")
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    text = text.strip()
                    if text:
                        blocks.append({
                            "text": text,
                            "sourceRef": f"Page {i + 1}"
                        })

        # 2. DOCX Parsing
        elif fn_lower.endswith('.docx') or 'wordprocessingml' in content_type or 'docx' in content_type:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            para_group = []
            group_counter = 1

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    para_group.append(p_text)
                    if len("\n".join(para_group)) >= 500:
                        blocks.append({
                            "text": "\n".join(para_group),
                            "sourceRef": f"Section {group_counter}"
                        })
                        para_group = []
                        group_counter += 1

            if para_group:
                blocks.append({
                    "text": "\n".join(para_group),
                    "sourceRef": f"Section {group_counter}"
                })

        # 3. PPTX Parsing
        elif fn_lower.endswith('.pptx') or 'presentationml' in content_type or 'pptx' in content_type:
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            for i, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        txt = shape.text.strip()
                        if txt:
                            slide_texts.append(txt)
                if slide_texts:
                    blocks.append({
                        "text": "\n".join(slide_texts),
                        "sourceRef": f"Slide {i + 1}"
                    })

        else:
            raise ParsingError(f"Unsupported file format '{filename}'. Allowed formats: .pdf, .docx, .pptx.")

    except ParsingError:
        raise
    except Exception as e:
        logger.error(f"Error parsing file '{filename}': {e}")
        raise ParsingError(f"Failed to parse document '{filename}': {str(e)}")

    if not blocks:
        raise ParsingError(f"No extractable text was found inside '{filename}'. Please ensure it is not a scanned image-only PDF.")

    return blocks
